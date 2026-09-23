"""Generate DOCX bytes locally without storing meeting content on disk."""
from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt

from backend.schemas import MeetingProtocol

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def export_protocol_docx(protocol: MeetingProtocol) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    document.core_properties.author = ""
    document.core_properties.title = "Протокол совещания"
    document.add_heading("Протокол совещания", 0)
    if protocol.title:
        document.add_paragraph(protocol.title, "Subtitle")
    document.add_heading("Summary", 1)
    document.add_paragraph(protocol.summary)
    document.add_heading("Поручения", 1)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.autofit = False
    widths = (7, 4.5, 3.5, 2)
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    for cell, label, width in zip(table.rows[0].cells, ("Поручение", "Ответственный", "Срок", "Confidence"), widths):
        cell.width = Cm(width)
        cell.text = label
        cell.paragraphs[0].runs[0].bold = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for item in protocol.action_items:
        values = (item.text, item.responsible or "Не определено", item.deadline or "Не определено",
                  "Не определено" if item.confidence is None else f"{item.confidence:.2f}")
        for cell, value, width in zip(table.add_row().cells, values, widths):
            cell.width = Cm(width)
            cell.text = value
    if not protocol.action_items:
        document.add_paragraph("Поручения не выявлены.")
    document.add_heading("Transcript", 1)
    for segment in protocol.transcript:
        paragraph = document.add_paragraph()
        if segment.start is not None or segment.end is not None:
            start = "?" if segment.start is None else f"{segment.start:g}"
            end = "?" if segment.end is None else f"{segment.end:g}"
            paragraph.add_run(f"[{start}–{end} с] ")
        if segment.speaker:
            paragraph.add_run(f"{segment.speaker}: ").bold = True
        paragraph.add_run(segment.text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
