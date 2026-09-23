"""Opt-in real local browser check; never runs during unittest discovery.

Usage: .venv-integration/bin/python tests/smoke_frontend.py '/path/to/meeting.mp3'
Requires a running localhost:8000 server, Playwright and installed Google Chrome.
Only metadata is printed. Meeting results stay in ignored audio_module/results.
"""
import argparse
from collections import Counter
from io import BytesIO
import json
from pathlib import Path
import time
from urllib.parse import urlparse

from docx import Document
from playwright.sync_api import sync_playwright


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('mp3', type=Path)
    args = parser.parse_args()
    source = args.mp3.expanduser().resolve()
    assert source.is_file(), 'MP3 not found'
    output = Path(__file__).resolve().parents[1] / 'audio_module' / 'results'
    output.mkdir(exist_ok=True)
    external, errors, calls = [], [], []
    with sync_playwright() as p:
        browser = p.chromium.launch(channel='chrome', headless=True,
            args=['--disable-background-networking', '--disable-component-update', '--disable-sync'])
        context = browser.new_context(accept_downloads=True, viewport={'width': 1440, 'height': 1000},
                                      service_workers='block')
        def local_only(route):
            url = urlparse(route.request.url)
            if url.scheme in ('http', 'https') and url.netloc != '127.0.0.1:8000':
                external.append(url.netloc)
                route.abort()
            else:
                calls.append((route.request.method, url.path))
                route.continue_()
        context.route('**/*', local_only)
        page = context.new_page()
        page.on('pageerror', lambda error: errors.append(type(error).__name__))
        health = context.request.get('http://127.0.0.1:8000/health')
        assert health.status == 200 and health.json() == {'status': 'ok'}
        assert page.goto('http://127.0.0.1:8000/').status == 200
        page.wait_for_function("document.getElementById('connection').textContent === 'Backend подключён'")
        page.locator('#audio-file').set_input_files(str(source))
        assert page.locator('#audio-limit').input_value() == '60'
        started = time.monotonic()
        with page.expect_response(lambda r: '/analyze-audio' in r.url and r.request.method == 'POST', timeout=600_000) as pending:
            page.locator('#process-button').click()
            assert page.locator('#process-button').is_disabled()
            assert page.locator('#audio-file').is_disabled()
            assert page.locator('#status-title').inner_text() == 'Обрабатываем аудио…'
            # A second submit is ignored even if triggered directly by the form.
            page.locator('#upload-form').evaluate("el => el.dispatchEvent(new Event('submit', {cancelable:true}))")
        response = pending.value
        assert response.status == 200, f'HTTP {response.status}'
        result = response.json()
        elapsed = round(time.monotonic() - started, 3)
        page.wait_for_function("document.getElementById('status-title').textContent === 'Готово'")
        assert not page.locator('#process-button').is_disabled()
        assert not page.locator('#download-button').is_disabled()
        assert result['diarization']['status'] == 'ok'
        assert result['diarization_available'] is True
        segments = result['protocol']['transcript']
        assert all(set(('text','start','end','speaker')) <= set(s) for s in segments)
        assert result['processing']['processed_start_seconds'] == 0
        assert result['processing']['processed_end_seconds'] == 60
        assert page.locator('.transcript-segment').count() == len(segments)
        assert page.locator('#task-list tr').count() == len(result['protocol']['action_items'])
        assert page.locator('#summary-list').inner_text().strip()
        assert page.locator('.results-heading + .processing-note').is_visible()
        for i, segment in enumerate(segments):
            row = page.locator('.transcript-segment').nth(i)
            assert row.locator('.transcript-text').inner_text() == segment['text']
            if segment['speaker']:
                assert row.locator('.speaker').inner_text() == segment['speaker']
            else:
                assert row.locator('small').is_visible()
        for width in (1440, 390):
            page.set_viewport_size({'width': width, 'height': 1000})
            assert page.evaluate('document.documentElement.scrollWidth <= window.innerWidth'), 'Page overflows'
        with page.expect_download(timeout=30_000) as pending_download:
            page.locator('#download-button').click()
        download = pending_download.value
        target = output / 'integration-browser-protocol.docx'
        download.save_as(target)
        page.wait_for_function("document.getElementById('document-note').textContent === 'DOCX получен с сервера.'")
        doc = Document(BytesIO(target.read_bytes()))
        paragraphs = '\n'.join(p.text for p in doc.paragraphs)
        assert result['protocol']['summary'] in paragraphs
        assert all(s['text'] in paragraphs for s in segments)
        assert len(doc.tables[0].rows) == len(result['protocol']['action_items']) + 1
        assert not external and not errors
        assert calls.count(('POST', '/analyze-audio')) == 1
        assert calls.count(('POST', '/export-docx')) == 1
        assert not any('/api/meetings' in path for _, path in calls)
        summary = {'health': 200, 'audio_http': response.status, 'status': result['diarization']['status'],
            'interval': [0, 60], 'language': result['language'], 'segments': len(segments),
            'speakers': result['diarization']['speakers'],
            'assigned': sum(s['speaker'] is not None for s in segments),
            'reasons': dict(Counter(a['reason'] for a in result['diarization']['assignments'])),
            'tasks': len(result['protocol']['action_items']), 'request_seconds': elapsed,
            'transcription_seconds': result['processing']['elapsed_seconds'],
            'diarization_seconds': result['diarization']['elapsed_seconds'],
            'docx_bytes': target.stat().st_size, 'docx_content_checked': True,
            'upload_requests': 1, 'export_requests': 1, 'external_browser_requests': len(external),
            'javascript_errors': len(errors), 'mobile_overflow': False}
        (output / 'integration-browser-result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2))
        (output / 'integration-browser-summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2))
        print(json.dumps(summary, ensure_ascii=False, indent=2))
        context.close()
        browser.close()


if __name__ == '__main__':
    main()
