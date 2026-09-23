"""Acceptance tests. Fake audio adapters below are test-only dependency overrides."""
from io import BytesIO
from pathlib import Path
import os
import subprocess
import sys
import unittest
from unittest.mock import patch
from zipfile import ZipFile
from docx import Document
from fastapi.testclient import TestClient
from pydantic import ValidationError
from backend.main import app
from backend.schemas import ActionItem, AnalyzeRequest, MeetingProtocol, TranscriptionResult
from backend.services.audio import InvalidAudioError, LocalWhisperTranscriber, NoSpeechError, get_transcriber
from backend.services.protocol_analyzer import analyze_transcript

SAMPLE = Path(__file__).with_name('sample_transcript.txt').read_text(encoding='utf-8')

class AnalysisTests(unittest.TestCase):
    def tasks(self, text): return analyze_transcript(text)['action_items']

    def test_realistic_three_tasks(self):
        result = analyze_transcript(SAMPLE)
        self.assertEqual([(i['text'],i['responsible'],i['deadline']) for i in result['action_items']], [
            ('Подготовить стратегию закупа сырья','Гульмира Сериковна','15 октября'),
            ('Подготовить финансовое решение','Тимур Болатович','30 сентября'),
            ('Проверить договор с подрядчиком','Айнур Каировна',None)])
        self.assertTrue(result['summary'])
        for item in result['action_items']:
            self.assertIn(item['source_fragment'],SAMPLE)
            self.assertGreaterEqual(item['confidence'],0)
            self.assertLessEqual(item['confidence'],1)

    def test_unknown_owner_and_deadline(self):
        item = self.tasks('Проверьте договор с подрядчиком.')[0]
        self.assertIsNone(item['responsible'])
        self.assertIsNone(item['deadline'])

    def test_speaker_is_not_assignee(self):
        result = analyze_transcript([{'speaker':'Асхат Ерланович','start':0.0,'end':5.0,'text':'Проверьте договор.'}])
        self.assertIsNone(result['action_items'][0]['responsible'])
        self.assertEqual(result['transcript'][0]['speaker'],'Асхат Ерланович')
        self.assertEqual(result['transcript'][0]['end'],5.0)

    def test_explicit_owner_department(self):
        item = self.tasks('Провести юридическую проверку договора — ответственный Юридический департамент, срок до 30 сентября.')[0]
        self.assertEqual(item['responsible'],'Юридический департамент')
        self.assertEqual(item['deadline'],'30 сентября')
        self.assertEqual(item['text'],'Провести юридическую проверку договора')

    def test_relative_deadlines_no_invented_dates(self):
        for phrase in ['до пятницы','на следующей неделе','до конца недели','за две недели','через три дня','завтра']:
            with self.subTest(phrase=phrase):
                item = self.tasks(f'Айнур Каировна, подготовьте отчет {phrase}.')[0]
                self.assertEqual(item['deadline'],phrase.removeprefix('до '))
                self.assertNotIn('2026',item['deadline'])

    def test_numeric_date(self):
        item = self.tasks('Иван, представьте отчет до 15.10.2026.')[0]
        self.assertEqual(item['deadline'],'15.10.2026')
        self.assertEqual(item['text'],'Представить отчет')

    def test_preposed_deadline(self):
        item = self.tasks('Айнур Каировна, до пятницы проверьте договор.')[0]
        self.assertEqual(item['deadline'],'пятницы')
        self.assertEqual(item['text'],'Проверить договор')

    def test_no_tasks_for_greeting(self):
        self.assertEqual(self.tasks('Коллеги, добрый день. Начинаем совещание.'),[])

    def test_ambiguous_and_negative_are_skipped(self):
        for text in ['Не отправьте письмо.', 'Не нужно подготовить отчет.',
                     'Если согласуют бюджет, подготовьте отчет.', 'Вы можете подготовить отчет?',
                     'Предлагаю подготовить отчет.', 'Вчера попросили подготовить отчет.',
                     'Цель проекта — подготовить отчет.']:
            with self.subTest(text=text): self.assertEqual(self.tasks(text),[])

    def test_group_is_not_person(self):
        self.assertIsNone(self.tasks('Коллеги, подготовьте отчет.')[0]['responsible'])

    def test_kazakh_limited_template(self):
        item = self.tasks('Айгүл, есепті 15 қазанға дейін дайындаңыз.')[0]
        self.assertEqual(item['responsible'],'Айгүл')
        self.assertEqual(item['deadline'],'15 қазанға дейін')
        self.assertIn('есепті',item['text'].lower())

    def test_mixed_limited_template(self):
        item = self.tasks('Айгүл, подготовьте есеп на следующей неделе.')[0]
        self.assertEqual(item['responsible'],'Айгүл')
        self.assertEqual(item['deadline'],'на следующей неделе')

    def test_no_cross_turn_inference(self):
        result = analyze_transcript([{'speaker':'Руководитель','text':'Айнур Каировна, вы курируете договор?'},
                                     {'speaker':'Айнур Каировна','text':'Да.'},
                                     {'speaker':'Руководитель','text':'Проверьте договор до пятницы.'}])
        self.assertIsNone(result['action_items'][0]['responsible'])

    def test_different_deadlines_are_not_silently_merged(self):
        self.assertEqual(len(self.tasks('Иван, подготовьте отчет до пятницы. Иван, подготовьте отчет до среды.')),2)

    def test_duplicate_identical_task(self):
        self.assertEqual(len(self.tasks('Иван, проверьте отчет. Иван, проверьте отчет.')),1)

    def test_summary_extracts_source_facts(self):
        result = analyze_transcript('Коллеги, начинаем совещание. Выпуск составляет 71 процент. Нужно проверить отчет.')
        self.assertIn('Выпуск составляет 71 процент.',result['summary'])

    def test_asr_sentence_split_across_chunks(self):
        result=analyze_transcript([{'text':'Подготовить отчет — ответственный'},
                                  {'text':'Иван Петрович, срок до 15 октября.'}])
        item=result['action_items'][0]
        self.assertEqual(item['responsible'],'Иван Петрович')
        self.assertEqual(item['deadline'],'15 октября')
        self.assertEqual(item['text'],'Подготовить отчет')

    def test_numbering_is_not_a_name(self):
        for label in ['Первое','Четвертое','Четвёртое','Пятая','Пятое']:
            with self.subTest(label=label):
                item=self.tasks(label+', подготовить отчет.')[0]
                self.assertIsNone(item['responsible'])

    def test_explicit_unknown_owner(self):
        for value in ['не указан','не определён','не назначен','неизвестно']:
            with self.subTest(value=value):
                item=self.tasks('Подготовить отчет — ответственный '+value+'.')[0]
                self.assertIsNone(item['responsible'])

    def test_sentence_chunks_do_not_cross_speakers(self):
        result=analyze_transcript([{'speaker':'Speaker 1','text':'Иван, проверьте отчет'},
                                  {'speaker':'Speaker 2','text':'до пятницы работает офис.'}])
        self.assertIsNone(result['action_items'][0]['deadline'])

    def test_deadline_word_inside_task_is_preserved(self):
        item=self.tasks('Пропишите в договоре срок выставления счета.')[0]
        self.assertIn('срок выставления счета',item['text'])
        self.assertIsNone(item['deadline'])

    def test_enumerated_task_with_en_dash(self):
        item=self.tasks('Первое – разработать стратегию закупа — ответственный Гульмира Сериковна, срок до 15 октября.')[0]
        self.assertEqual(item['responsible'],'Гульмира Сериковна')
        self.assertEqual(item['deadline'],'15 октября')

    def test_adjacent_explicit_metadata(self):
        items=analyze_transcript([{'text':'Подготовить отчет.'},
                                  {'text':'Ответственный Тимур Болатович.'},
                                  {'text':'Срок до 30 сентября.'}])['action_items']
        self.assertEqual(len(items),1)
        self.assertEqual(items[0]['responsible'],'Тимур Болатович')
        self.assertEqual(items[0]['deadline'],'30 сентября')
        self.assertIn('Ответственный Тимур Болатович.',items[0]['source_fragment'])

    def test_metadata_cannot_cross_speaker(self):
        result=analyze_transcript([{'speaker':'Speaker 1','text':'Подготовить отчет.'},
                                  {'speaker':'Speaker 2','text':'Ответственный Иван.'}])
        self.assertIsNone(result['action_items'][0]['responsible'])

class ValidationTests(unittest.TestCase):
    def test_invalid_inputs(self):
        cases = [{},{'text':' '},{'text':42},{'text':'a','transcript':[{'text':'b'}]},
                 {'transcript':[]},{'text':'a','language':'ru'},
                 {'transcript':[{'text':'a','start':5.0,'end':4.0}]},
                 {'transcript':[{'text':'a','start':-1.0}]},
                 {'transcript':[{'text':'a','start':float('inf')}]},
                 {'text':'a'*100001},{'transcript':[{'text':'a'}]*2001},
                 {'text':'a\n'*2001},{'text':'x\x00y'},
                 {'transcript':[{'text':'a'*60000}]*2}]
        for data in cases:
            with self.subTest(keys=list(data)):
                with self.assertRaises(ValidationError): AnalyzeRequest.model_validate(data)

    def test_confidence_bounds(self):
        for value in [-0.01,1.01,float('nan'),float('inf')]:
            with self.assertRaises(ValidationError): ActionItem(text='a',source_fragment='a',confidence=value)

    def test_model_defaults_are_null(self):
        item = ActionItem(text='Проверить отчет',source_fragment='Проверьте отчет.')
        self.assertIsNone(item.responsible)
        self.assertIsNone(item.deadline)
        self.assertIsNone(item.confidence)

class APITests(unittest.TestCase):
    def setUp(self): self.client = TestClient(app,raise_server_exceptions=False)
    def tearDown(self):
        app.dependency_overrides.clear()
        self.client.close()

    def test_health(self):
        response = self.client.get('/health')
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.json(),{'status':'ok'})

    def test_analysis(self):
        response = self.client.post('/analyze',json={'text':SAMPLE})
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(len(response.json()['action_items']),3)

    def test_structured_input(self):
        data = {'title':'Тест','transcript':[{'speaker':'Speaker 1','start':0.0,'end':5.0,'text':'Проверьте договор.'}]}
        response = self.client.post('/analyze',json=data)
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(response.json()['transcript'],data['transcript'])

    def test_empty_speech(self):
        self.assertEqual(self.client.post('/analyze',json={'text':'Speaker 1:'}).status_code,422)

    def test_validation_hides_meeting_text(self):
        response = self.client.post('/analyze',json={'text':'PRIVATE MEETING','extra':'SECRET'})
        self.assertEqual(response.status_code,422)
        self.assertNotIn('PRIVATE MEETING',response.text)
        self.assertNotIn('SECRET',response.text)
        self.assertTrue(all(set(e)<= {'loc','msg','type'} for e in response.json()['detail']))

    def test_internal_value_error_is_500(self):
        with patch('backend.main.build_protocol',side_effect=ValueError('PRIVATE')):
            response = self.client.post('/analyze',json={'text':'Проверьте отчет.'})
        self.assertEqual(response.status_code,500)
        self.assertNotIn('PRIVATE',response.text)
        self.assertNotIn('не содержит',response.text)

    def test_export_real_analysis_result(self):
        protocol = self.client.post('/analyze',json={'text':SAMPLE}).json()
        response = self.client.post('/export-docx',json=protocol)
        self.assertEqual(response.status_code,200,response.text[:100])
        self.assertEqual(response.headers['content-type'],'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        self.assertIn('meeting_protocol.docx',response.headers['content-disposition'])
        self.assertIsNone(ZipFile(BytesIO(response.content)).testzip())
        doc = Document(BytesIO(response.content))
        self.assertEqual(len(doc.tables[0].rows),4)
        self.assertEqual(doc.tables[0].rows[-1].cells[2].text,'Не определено')
        self.assertIn(protocol['summary'],[p.text for p in doc.paragraphs])
        self.assertTrue(all(any(s['text'] in p.text for p in doc.paragraphs) for s in protocol['transcript']))

    def test_export_no_tasks_and_nulls(self):
        protocol = {'summary':'Обсуждение','transcript':[],'action_items':[]}
        response = self.client.post('/export-docx',json=protocol)
        self.assertEqual(response.status_code,200)
        self.assertIn('Поручения не выявлены.',[p.text for p in Document(BytesIO(response.content)).paragraphs])
        protocol['action_items']=[{'text':'Проверить отчет','source_fragment':'Проверьте отчет.'}]
        response = self.client.post('/export-docx',json=protocol)
        self.assertEqual([c.text for c in Document(BytesIO(response.content)).tables[0].rows[1].cells][1:],['Не определено']*3)

    def test_export_rejects_bad_protocol(self):
        self.assertEqual(self.client.post('/export-docx',json={}).status_code,422)

    def test_json_body_limit(self):
        response = self.client.post('/analyze',content=b'x'*(2*1024*1024+1),headers={'Content-Type':'application/json'})
        self.assertEqual(response.status_code,413)

    def test_openapi_documents_audio(self):
        spec = self.client.get('/openapi.json').json()
        self.assertIn('/transcribe',spec['paths'])
        self.assertIn('multipart/form-data',spec['paths']['/analyze-audio']['post']['requestBody']['content'])

    def test_audio_not_configured_is_503(self):
        app.dependency_overrides[get_transcriber]=lambda:LocalWhisperTranscriber()
        with patch.dict(os.environ,{'WHISPER_MODEL_DIR':'/missing-local-model','WHISPER_MODEL_PATH':''}):
            response = self.client.post('/transcribe',files={'file':('test.mp3',b'audio','audio/mpeg')})
        self.assertEqual(response.status_code,503)
        self.assertIn('WHISPER_MODEL_DIR',response.text)

    def test_audio_file_validation(self):
        self.assertEqual(self.client.post('/transcribe',files={'file':('a.txt',b'x','audio/mpeg')}).status_code,415)
        self.assertEqual(self.client.post('/transcribe',files={'file':('a.mp3',b'')}).status_code,422)
        self.assertEqual(self.client.post('/transcribe',files={'file':('a.mpeg',b'','audio/mpeg')}).status_code,422)
        self.assertEqual(self.client.post('/transcribe').status_code,422)
        self.assertEqual(self.client.post('/transcribe?language=xx',files={'file':('a.mp3',b'x')}).status_code,422)
        with patch('backend.main.MAX_UPLOAD_BYTES',4):
            self.assertEqual(self.client.post('/transcribe',files={'file':('a.mp3',b'12345')}).status_code,413)

    def test_audio_adapter_contract_and_cleanup(self):
        paths=[]
        class FakeAudio:
            def transcribe(self,path,*,language=None):
                paths.append(path)
                assert path.read_bytes()==b'test audio bytes'
                assert path.name=='upload.mp3'
                assert language=='ru'
                return TranscriptionResult(text='Айнур Каировна, проверьте договор.',
                    transcript=[{'speaker':'Speaker 2','start':1.0,'end':3.0,'text':'Айнур Каировна, проверьте договор.'}],
                    language='ru',diarization_available=True,diarization={'status':'ok'},warnings=[])
        app.dependency_overrides[get_transcriber]=lambda:FakeAudio()
        for route in ('/transcribe','/analyze-audio'):
            for name in ('../../evil.mp3','meeting.mpeg','MEETING.MPEG'):
                with self.subTest(route=route,filename=name):
                    response = self.client.post(route+'?language=ru',
                        files={'file':(name,b'test audio bytes','audio/mpeg')})
                    self.assertEqual(response.status_code,200,response.text)
                    self.assertTrue(response.json()['diarization_available'])
                    if route=='/analyze-audio':
                        protocol=response.json()['protocol']
                        self.assertEqual(protocol['action_items'][0]['responsible'],'Айнур Каировна')
                        self.assertEqual(protocol['transcript'][0]['speaker'],'Speaker 2')
                        self.assertEqual(self.client.post('/export-docx',json=protocol).status_code,200)
                    else:
                        self.assertEqual(response.json()['transcript'][0]['speaker'],'Speaker 2')
        self.assertTrue(all(not p.exists() for p in paths))

    def test_audio_error_and_cleanup(self):
        for error,status in [(InvalidAudioError('Некорректное аудио'),422),(NoSpeechError(),422),(RuntimeError('PRIVATE'),500)]:
            paths=[]
            class FailingAudio:
                def transcribe(self,path,*,language=None):
                    paths.append(path)
                    raise error
            app.dependency_overrides[get_transcriber]=lambda:FailingAudio()
            response=self.client.post('/transcribe',files={'file':('a.mp3',b'x')})
            self.assertEqual(response.status_code,status)
            self.assertNotIn('PRIVATE',response.text)
            self.assertTrue(all(not p.exists() for p in paths))

    def test_blank_audio_title_is_validation_error(self):
        response=self.client.post('/analyze-audio?title=%20%20',files={'file':('a.mp3',b'x')})
        self.assertEqual(response.status_code,422)

    def test_audio_result_cannot_claim_diarization_without_ok_status(self):
        with self.assertRaises(ValidationError):
            TranscriptionResult(text='Привет',transcript=[{'text':'Привет'}],diarization_available=True)

    def test_body_limit_without_content_length(self):
        chunks=(b'x'*1048576 for _ in range(3))
        response=self.client.post('/analyze',content=chunks,headers={'Content-Type':'application/json'})
        self.assertEqual(response.status_code,413)

    def test_frontend_cors_is_explicit(self):
        code="""
import os
os.environ['FRONTEND_ORIGINS']='http://localhost:5173'
from backend.main import app
from fastapi.testclient import TestClient
with TestClient(app) as client:
    good=client.options('/analyze',headers={'Origin':'http://localhost:5173','Access-Control-Request-Method':'POST','Access-Control-Request-Headers':'content-type'})
    assert good.status_code==200 and good.headers['access-control-allow-origin']=='http://localhost:5173'
    bad=client.options('/analyze',headers={'Origin':'https://untrusted.example','Access-Control-Request-Method':'POST'})
    assert bad.status_code==400
"""
        result=subprocess.run([sys.executable,'-c',code],capture_output=True,text=True)
        self.assertEqual(result.returncode,0,result.stderr)

class OfflineTests(unittest.TestCase):
    def test_text_and_docx_need_no_network(self):
        with patch('socket.socket.connect',side_effect=AssertionError('Network forbidden')):
            result=analyze_transcript(SAMPLE)
            from backend.services.docx_exporter import export_protocol_docx
            self.assertGreater(len(export_protocol_docx(MeetingProtocol.model_validate(result))),1000)

    def test_import_no_audio_dependency_or_ai_package(self):
        code="""
import importlib.abc, sys
class Block(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path, target=None):
        if fullname.split('.')[0] in {'ai','faster_whisper','av'}:
            raise ImportError('Optional module blocked: '+fullname)
sys.meta_path.insert(0, Block())
from backend.main import app
print(app.title)
"""
        result=subprocess.run([sys.executable,'-c',code],capture_output=True,text=True)
        self.assertEqual(result.returncode,0,result.stderr)

if __name__=='__main__': unittest.main()
