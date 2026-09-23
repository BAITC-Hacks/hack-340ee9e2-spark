"""Offline API regression tests; no model or audio dependencies required."""
from io import BytesIO
from pathlib import Path
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient

from backend.main import app
from backend.services.protocol_analyzer import analyze_transcript
from backend.services.docx_exporter import DOCX_MEDIA_TYPE


class BackendServicesTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)
        self.addCleanup(self.client.close)

    def test_health(self):
        response = self.client.get("/health")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {"status": "ok"})

    def test_sample_and_export(self):
        text = Path(__file__).with_name("sample_transcript.txt").read_text(encoding="utf-8")
        response = self.client.post("/analyze", json={"text": text})
        self.assertEqual(response.status_code, 200)
        protocol = response.json()
        self.assertEqual(len(protocol["action_items"]), 3)
        self.assertEqual([a["text"] for a in protocol["action_items"]], [
            "Подготовить стратегию закупа сырья", "Подготовить финансовое решение",
            "Проверить договор с подрядчиком",
        ])
        self.assertEqual([(a["responsible"], a["deadline"]) for a in protocol["action_items"]],
                         [("Гульмира Сериковна", "15 октября"), ("Тимур Болатович", "30 сентября"),
                          ("Айнур Каировна", None)])
        for item in protocol["action_items"]:
            self.assertIn(item["source_fragment"], text)
        exported = self.client.post("/export-docx", json=protocol)
        self.assertEqual(exported.status_code, 200)
        self.assertEqual(exported.headers["content-type"], DOCX_MEDIA_TYPE)
        self.assertIn("attachment", exported.headers["content-disposition"])
        self.assertGreater(len(exported.content), 0)
        with ZipFile(BytesIO(exported.content)) as archive:
            self.assertIsNone(archive.testzip())
            self.assertIn("word/document.xml", archive.namelist())
        doc = Document(BytesIO(exported.content))
        self.assertEqual(doc.paragraphs[0].text, "Meeting Protocol")
        self.assertEqual(len(doc.tables[0].rows), 4)
        self.assertEqual(doc.tables[0].cell(1, 1).text, "Гульмира Сериковна")
        self.assertEqual(doc.tables[0].cell(3, 2).text, "Не определено")
        paragraphs = [p.text for p in doc.paragraphs]
        for heading in ("1. Summary", "2. Action Items", "3. Transcript"):
            self.assertIn(heading, paragraphs)
        self.assertIn(protocol["summary"], paragraphs)
        self.assertTrue(any("Айнур Каировна, проверьте договор с подрядчиком." in p for p in paragraphs))

    def test_segments_preserve_metadata_and_unknowns(self):
        segments = [{"speaker": None, "start": 1.0, "end": 3.5, "text": "Нужно подготовить отчёт."}]
        response = self.client.post("/analyze", json={"transcript": segments, "title": "Тест"})
        self.assertEqual(response.status_code, 200)
        result = response.json()
        self.assertEqual(result["transcript"], segments)
        self.assertEqual(result["title"], "Тест")
        self.assertIsNone(result["action_items"][0]["responsible"])
        self.assertIsNone(result["action_items"][0]["deadline"])
        exported = self.client.post("/export-docx", json=result)
        doc = Document(BytesIO(exported.content))
        self.assertEqual(doc.tables[0].cell(1, 1).text, "Не определено")
        self.assertEqual(doc.tables[0].cell(1, 2).text, "Не определено")

    def test_analysis_and_export_do_not_open_network_connections(self):
        with patch("socket.socket.connect", side_effect=AssertionError("Network access forbidden")):
            protocol = self.client.post("/analyze", json={"text": "Нужно проверить договор."})
            self.assertEqual(protocol.status_code, 200)
            self.assertEqual(self.client.post("/export-docx", json=protocol.json()).status_code, 200)

    def test_explicit_speaker_and_relative_deadline(self):
        result = analyze_transcript("Айнур Каировна: Проверьте договор до пятницы.")
        self.assertIsNone(result["action_items"][0]["responsible"])
        self.assertEqual(result["action_items"][0]["deadline"], "пятницы")

    def test_export_preserves_zero_and_unknown_confidence(self):
        protocol = analyze_transcript("Нужно проверить договор.")
        for confidence, expected in ((0.0, "0.00"), (None, "Не определено")):
            protocol["action_items"][0]["confidence"] = confidence
            response = self.client.post("/export-docx", json=protocol)
            self.assertEqual(response.status_code, 200)
            doc = Document(BytesIO(response.content))
            self.assertEqual(doc.tables[0].cell(1, 3).text, expected)

    def test_invalid_export_and_timestamps(self):
        self.assertEqual(self.client.post("/export-docx", json={}).status_code, 422)
        response = self.client.post("/analyze", json={"transcript": [
            {"text": "Тест", "start": 5.0, "end": 1.0},
        ]})
        self.assertEqual(response.status_code, 422)

    def test_negated_and_uncertain_tasks_ignored(self):
        result = analyze_transcript("Не нужно подготовить отчёт. Возможно, проведите встречу? Обсудили новости.")
        self.assertEqual(result["action_items"], [])
        self.assertEqual(self.client.post("/export-docx", json=result).status_code, 200)

    def test_invalid_requests(self):
        for payload in ({}, {"text": " "}, {"text": "Асхат Ерланович:"},
                        {"text": "Тест", "transcript": [{"text": "Тест"}]}):
            self.assertEqual(self.client.post("/analyze", json=payload).status_code, 422)


if __name__ == "__main__":
    unittest.main()
